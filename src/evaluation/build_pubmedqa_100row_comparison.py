from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_DIR = Path(__file__).resolve().parents[2]

DATASET_FILE = (
    PROJECT_DIR
    / "data"
    / "evaluation"
    / "pubmedqa_100row_evaluation_set.csv"
)

BASE_FILE = (
    PROJECT_DIR
    / "results"
    / "pubmedqa_100row_comparison"
    / "base_qwen_100row_predictions.csv"
)

LORA_FILE = (
    PROJECT_DIR
    / "results"
    / "pubmedqa_100row_comparison"
    / "pubmedqa_lora_100row_predictions.csv"
)

OUTPUT_DIR = (
    PROJECT_DIR
    / "results"
    / "pubmedqa_100row_comparison"
)

OUTPUT_CSV = (
    OUTPUT_DIR
    / "pubmedqa_100row_model_comparison.csv"
)

OUTPUT_XLSX = (
    OUTPUT_DIR
    / "pubmedqa_100row_model_comparison.xlsx"
)


# ------------------------------------------------------------
# Load files
# ------------------------------------------------------------

dataset = pd.read_csv(
    DATASET_FILE,
    dtype={"pubmed_id": str},
)

base = pd.read_csv(
    BASE_FILE,
    dtype={"pubmed_id": str},
)

lora = pd.read_csv(
    LORA_FILE,
    dtype={"pubmed_id": str},
)


# ------------------------------------------------------------
# Keep only the useful prediction columns
# ------------------------------------------------------------

base = base[
    [
        "sample_id",
        "pubmed_id",
        "base_qwen_prediction",
        "correct",
    ]
].rename(
    columns={
        "correct": "base_qwen_correct",
    }
)

lora = lora[
    [
        "sample_id",
        "pubmed_id",
        "pubmedqa_lora_prediction",
        "correct",
    ]
].rename(
    columns={
        "correct": "pubmedqa_lora_correct",
    }
)


# ------------------------------------------------------------
# Merge everything row-by-row
# ------------------------------------------------------------

comparison = dataset.merge(
    base,
    on=[
        "sample_id",
        "pubmed_id",
    ],
    how="left",
)

comparison = comparison.merge(
    lora,
    on=[
        "sample_id",
        "pubmed_id",
    ],
    how="left",
)


# ------------------------------------------------------------
# Add simple comparison columns
# ------------------------------------------------------------

comparison["base_result"] = comparison[
    "base_qwen_correct"
].map(
    {
        True: "Correct",
        False: "Wrong",
    }
)

comparison["lora_result"] = comparison[
    "pubmedqa_lora_correct"
].map(
    {
        True: "Correct",
        False: "Wrong",
    }
)

comparison["improvement_case"] = ""

comparison.loc[
    (
        comparison["base_qwen_correct"] == False
    )
    &
    (
        comparison["pubmedqa_lora_correct"] == True
    ),
    "improvement_case",
] = "LoRA fixed Base-Qwen error"

comparison.loc[
    (
        comparison["base_qwen_correct"] == True
    )
    &
    (
        comparison["pubmedqa_lora_correct"] == True
    ),
    "improvement_case",
] = "Both correct"

comparison.loc[
    (
        comparison["base_qwen_correct"] == False
    )
    &
    (
        comparison["pubmedqa_lora_correct"] == False
    ),
    "improvement_case",
] = "Both wrong"

comparison.loc[
    (
        comparison["base_qwen_correct"] == True
    )
    &
    (
        comparison["pubmedqa_lora_correct"] == False
    ),
    "improvement_case",
] = "Base correct / LoRA wrong"


# ------------------------------------------------------------
# Summary values
# ------------------------------------------------------------

base_correct = int(
    comparison[
        "base_qwen_correct"
    ].sum()
)

lora_correct = int(
    comparison[
        "pubmedqa_lora_correct"
    ].sum()
)

base_accuracy = (
    base_correct
    / len(comparison)
)

lora_accuracy = (
    lora_correct
    / len(comparison)
)

fixed_errors = int(
    (
        comparison[
            "improvement_case"
        ]
        == "LoRA fixed Base-Qwen error"
    ).sum()
)

new_errors = int(
    (
        comparison[
            "improvement_case"
        ]
        == "Base correct / LoRA wrong"
    ).sum()
)


summary = pd.DataFrame(
    [
        {
            "Model":
                "Base Qwen2.5-3B-Instruct",

            "Rows":
                100,

            "Correct":
                base_correct,

            "Accuracy":
                base_accuracy,

            "Macro-F1":
                0.058,

            "YES Recall":
                0.04,

            "NO Recall":
                0.02,

            "Predicted YES":
                2,

            "Predicted NO":
                1,

            "Predicted MAYBE":
                97,
        },

        {
            "Model":
                "Qwen2.5-3B + PubMedQA LoRA",

            "Rows":
                100,

            "Correct":
                lora_correct,

            "Accuracy":
                lora_accuracy,

            "Macro-F1":
                0.819,

            "YES Recall":
                0.88,

            "NO Recall":
                0.76,

            "Predicted YES":
                56,

            "Predicted NO":
                44,

            "Predicted MAYBE":
                0,
        },
    ]
)


comparison_summary = pd.DataFrame(
    [
        [
            "Accuracy improvement",
            f"{(lora_accuracy - base_accuracy) * 100:.1f} percentage points",
        ],
        [
            "Base Qwen correct",
            f"{base_correct}/100",
        ],
        [
            "PubMedQA LoRA correct",
            f"{lora_correct}/100",
        ],
        [
            "Base errors corrected by LoRA",
            fixed_errors,
        ],
        [
            "Cases Base got right but LoRA missed",
            new_errors,
        ],
        [
            "Interpretation",
            (
                "LoRA substantially improved performance on the "
                "bounded PubMedQA evidence-conclusion task. "
                "This is not clinical accuracy."
            ),
        ],
    ],
    columns=[
        "Measure",
        "Value",
    ],
)


# ------------------------------------------------------------
# Save CSV
# ------------------------------------------------------------

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

comparison.to_csv(
    OUTPUT_CSV,
    index=False,
)


# ------------------------------------------------------------
# Save Excel workbook
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_XLSX,
    engine="openpyxl",
) as writer:

    summary.to_excel(
        writer,
        sheet_name="Model_Summary",
        index=False,
    )

    comparison_summary.to_excel(
        writer,
        sheet_name="Comparison_Summary",
        index=False,
    )

    comparison.to_excel(
        writer,
        sheet_name="Row_by_Row",
        index=False,
    )


# ------------------------------------------------------------
# Format workbook
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_XLSX
)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

header_font = Font(
    color="FFFFFF",
    bold=True,
)


for sheet_name in workbook.sheetnames:

    ws = workbook[
        sheet_name
    ]

    ws.freeze_panes = "A2"

    for cell in ws[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


# Summary formatting
summary_ws = workbook[
    "Model_Summary"
]

for col in range(
    1,
    summary_ws.max_column + 1,
):
    summary_ws.column_dimensions[
        chr(64 + col)
    ].width = 20

summary_ws.column_dimensions[
    "A"
].width = 38


# Comparison summary formatting
comparison_ws = workbook[
    "Comparison_Summary"
]

comparison_ws.column_dimensions[
    "A"
].width = 40

comparison_ws.column_dimensions[
    "B"
].width = 85


# Row-by-row formatting
row_ws = workbook[
    "Row_by_Row"
]

column_widths = {
    "A": 12,
    "B": 18,
    "C": 16,
    "D": 20,
    "E": 55,
    "F": 90,
    "G": 30,
    "H": 16,
    "I": 24,
    "J": 18,
    "K": 24,
    "L": 18,
    "M": 22,
    "N": 22,
    "O": 30,
}

for column, width in column_widths.items():

    row_ws.column_dimensions[
        column
    ].width = width


workbook.save(
    OUTPUT_XLSX
)


# ------------------------------------------------------------
# Print results
# ------------------------------------------------------------

print()
print("=" * 72)
print("PUBMEDQA 100-ROW BASE VS LORA COMPARISON")
print("=" * 72)

print()
print(
    f"Base Qwen:       "
    f"{base_correct}/100 correct "
    f"({base_accuracy * 100:.1f}%)"
)

print(
    f"PubMedQA LoRA:   "
    f"{lora_correct}/100 correct "
    f"({lora_accuracy * 100:.1f}%)"
)

print(
    f"Improvement:     "
    f"+{(lora_accuracy - base_accuracy) * 100:.1f} "
    f"percentage points"
)

print()

print(
    f"Base errors fixed by LoRA: "
    f"{fixed_errors}"
)

print(
    f"Base-correct cases lost by LoRA: "
    f"{new_errors}"
)

print()

print(
    "Comparison CSV:"
)
print(
    OUTPUT_CSV
)

print()

print(
    "Comparison Excel:"
)
print(
    OUTPUT_XLSX
)

print()

print(
    "Important: this is task-specific "
    "PubMedQA evidence-conclusion accuracy."
)

print(
    "It is not clinical or overall healthcare accuracy."
)

print("=" * 72)