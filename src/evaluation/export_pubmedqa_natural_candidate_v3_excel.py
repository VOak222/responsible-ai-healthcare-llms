from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_DIR = Path(__file__).resolve().parents[2]

INPUT_FILE = (
    PROJECT_DIR
    / "data"
    / "evaluation"
    / "pubmedqa_natural_candidate_test_v3_100.csv"
)

OUTPUT_FILE = (
    PROJECT_DIR
    / "data"
    / "evaluation"
    / "pubmedqa_natural_candidate_test_v3_100.xlsx"
)


df = pd.read_csv(
    INPUT_FILE,
    dtype={"pubmed_id": str},
)


readme = pd.DataFrame(
    {
        "Field": [
            "Dataset",
            "Purpose",
            "Questions",
            "Total candidates",
            "Supported",
            "Hallucinated",
            "Generation method",
            "Important limitation",
        ],
        "Description": [
            "PubMedQA-derived natural-language counterfactual evaluation set V3",
            "Cross-dataset evaluation of supported vs hallucinated candidate-answer classification",
            "50 unique PubMedQA research questions",
            "100 candidate answers",
            "50 evidence-aligned candidates",
            "50 counterfactual hallucinated candidates",
            (
                "Qwen2.5-3B-Instruct first generated an evidence-aligned "
                "candidate. A second generation step produced the opposite "
                "counterfactual answer without access to the original evidence."
            ),
            (
                "This is a constructed cross-dataset stress test, not an "
                "official PubMedQA or Med-HALT benchmark."
            ),
        ],
    }
)


with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl",
) as writer:

    df.to_excel(
        writer,
        sheet_name="Evaluation_Set",
        index=False,
    )

    readme.to_excel(
        writer,
        sheet_name="README",
        index=False,
    )

    workbook = writer.book

    evaluation_sheet = workbook[
        "Evaluation_Set"
    ]

    readme_sheet = workbook[
        "README"
    ]

    for cell in evaluation_sheet[1]:
        cell.font = Font(bold=True)

    for cell in readme_sheet[1]:
        cell.font = Font(bold=True)

    evaluation_sheet.freeze_panes = "A2"
    evaluation_sheet.auto_filter.ref = (
        evaluation_sheet.dimensions
    )

    widths = {
        "A": 16,
        "B": 18,
        "C": 16,
        "D": 18,
        "E": 55,
        "F": 70,
        "G": 70,
        "H": 20,
        "I": 24,
        "J": 28,
        "K": 18,
        "L": 14,
        "M": 16,
        "N": 28,
    }

    for column, width in widths.items():
        evaluation_sheet.column_dimensions[
            column
        ].width = width

    for row in evaluation_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    readme_sheet.column_dimensions[
        "A"
    ].width = 28

    readme_sheet.column_dimensions[
        "B"
    ].width = 100

    for row in readme_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


print()
print("=" * 72)
print("V3 EXCEL EXPORT COMPLETE")
print("=" * 72)
print()
print("Rows:", len(df))
print(
    "Supported:",
    (
        df["expected_medhalt_label"]
        == "supported"
    ).sum(),
)
print(
    "Hallucinated:",
    (
        df["expected_medhalt_label"]
        == "hallucinated"
    ).sum(),
)
print()
print("Output:")
print(OUTPUT_FILE)
print("=" * 72)