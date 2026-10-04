import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parents[2]

CSV_FILE = (
    PROJECT_DIR
    / "data"
    / "evaluation"
    / "pubmedqa_100row_evaluation_set.csv"
)

EXCEL_FILE = (
    PROJECT_DIR
    / "data"
    / "evaluation"
    / "pubmedqa_100row_evaluation_set.xlsx"
)


# ---------------------------------------------------------
# Load the exact CSV we already verified
# ---------------------------------------------------------

df = pd.read_csv(
    CSV_FILE,
    dtype={"pubmed_id": str},
)

if len(df) != 100:
    raise ValueError(
        f"Expected exactly 100 rows, but found {len(df)}."
    )

label_counts = df["ground_truth_label"].value_counts().to_dict()

if label_counts.get("yes", 0) != 50 or label_counts.get("no", 0) != 50:
    raise ValueError(
        f"Expected 50 YES and 50 NO rows. Found: {label_counts}"
    )


# ---------------------------------------------------------
# Export dataset
# ---------------------------------------------------------

with pd.ExcelWriter(
    EXCEL_FILE,
    engine="openpyxl",
) as writer:

    df.to_excel(
        writer,
        sheet_name="Evaluation_Set",
        index=False,
    )

    readme_data = [
        ["PUBMEDQA 100-ROW EVALUATION SET", ""],
        [
            "Purpose",
            "Compare Base Qwen2.5-3B-Instruct with "
            "Qwen2.5-3B + PubMedQA LoRA on the same test rows.",
        ],
        [
            "Source",
            "Existing frozen controlled PubMedQA test set.",
        ],
        [
            "Rows",
            "100 total: 50 YES and 50 NO.",
        ],
        [
            "Selection",
            "Fixed reproducible sample using random seed 42.",
        ],
        [
            "Important",
            "These rows are evaluation-only and must not be used "
            "for training, checkpoint selection, or threshold tuning.",
        ],
        [
            "Ground truth",
            "Original PubMedQA yes/no label.",
        ],
        [
            "Fair comparison",
            "Both Base Qwen and PubMedQA LoRA must receive exactly "
            "the same research question and PubMed abstract.",
        ],
    ]

    readme_df = pd.DataFrame(
        readme_data,
        columns=["Item", "Description"],
    )

    readme_df.to_excel(
        writer,
        sheet_name="README",
        index=False,
        header=False,
    )


# ---------------------------------------------------------
# Apply readable formatting
# ---------------------------------------------------------

workbook = load_workbook(EXCEL_FILE)

ws = workbook["Evaluation_Set"]

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

header_font = Font(
    color="FFFFFF",
    bold=True,
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


# Freeze column headers
ws.freeze_panes = "A2"


# Add filter
ws.auto_filter.ref = ws.dimensions


# Wrap all cells
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


# Column widths
column_widths = {
    "A": 12,   # sample_id
    "B": 18,   # source row
    "C": 16,   # PubMed ID
    "D": 20,   # label
    "E": 55,   # question
    "F": 90,   # evidence
    "G": 32,   # source split
    "H": 16,   # seed
}

for column, width in column_widths.items():
    ws.column_dimensions[column].width = width


# README formatting
readme_ws = workbook["README"]

readme_ws.column_dimensions["A"].width = 24
readme_ws.column_dimensions["B"].width = 100

for row in readme_ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

readme_ws["A1"].font = Font(
    bold=True,
    color="FFFFFF",
)

readme_ws["B1"].font = Font(
    bold=True,
    color="FFFFFF",
)

readme_ws["A1"].fill = header_fill
readme_ws["B1"].fill = header_fill


workbook.save(EXCEL_FILE)


print()
print("=" * 70)
print("PUBMEDQA 100-ROW EXCEL FILE CREATED")
print("=" * 70)
print(f"CSV source:  {CSV_FILE}")
print(f"Excel file:  {EXCEL_FILE}")
print()
print(f"Rows:        {len(df)}")
print(f"YES rows:    {label_counts.get('yes', 0)}")
print(f"NO rows:     {label_counts.get('no', 0)}")
print()
print("Sheets:")
print("  1. Evaluation_Set")
print("  2. README")
print("=" * 70)